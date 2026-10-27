import pptx

prs = pptx.Presentation('doc/NEXUS_Project_PPT.pptx')
for idx in range(5, 15):
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx + 1} ====================")
    title = slide.shapes.title.text if slide.shapes.title else "No Title"
    print(f"Title: {title}")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            text = "\n".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
            if text:
                print(f"  [Shape {s_idx} ({shape.name})]:")
                for line in text.split("\n")[:10]:
                    print(f"    {line}")
        elif shape.has_table:
            print(f"  [Table {s_idx}]: {len(shape.table.rows)} rows x {len(shape.table.columns)} cols")
            for r in shape.table.rows:
                row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
                print(f"    {row_vals}")
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            print(f"  [Picture {s_idx} ({shape.name})]")
