import pptx

prs = pptx.Presentation('doc/NEXUS_Project_PPT.pptx')

slides_to_show = [0, 4, 6, 11, 12, 13, 14, 17, 18]

for idx in slides_to_show:
    slide = prs.slides[idx]
    print(f"\n==================== SLIDE {idx + 1}: {slide.shapes.title.text if slide.shapes.title else 'No Title'} ====================")
    for s_idx, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            print(f"  Shape {s_idx} ({shape.name}):")
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                if p.text.strip():
                    print(f"    [P{p_idx}]: {p.text}")
        elif shape.has_table:
            print(f"  Table in Shape {s_idx} ({shape.name}):")
            for r_idx, row in enumerate(shape.table.rows):
                row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
                print(f"    [R{r_idx}]: {row_txt}")

