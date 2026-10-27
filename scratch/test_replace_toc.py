import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def build_perfect_report():
    doc = docx.Document('doc/DataMining_Report.docx')

    # Locate paragraph 105: "TABLE OF CONTENTS"
    toc_heading_idx = None
    for i, p in enumerate(doc.paragraphs):
        if 'TABLE OF CONTENTS' in p.text.upper():
            toc_heading_idx = i
            break
    
    print(f'Found TABLE OF CONTENTS at paragraph index: {toc_heading_idx}')
    p_toc_heading = doc.paragraphs[toc_heading_idx]

    # Ensure TABLE OF CONTENTS heading is centered, bold, 14pt Times New Roman
    p_toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_toc_heading.paragraph_format.space_before = Pt(12)
    p_toc_heading.paragraph_format.space_after = Pt(18)
    for r in p_toc_heading.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # TOC data: CHAPTER NO. | TITLE | PAGE NO.
    toc_data = [
        # (Chapter No, Title, Page No, is_main_chapter)
        ("", "ABSTRACT", "3", True),
        ("1", "INTRODUCTION", "4", True),
        ("1.1", "Problem Statement", "4", False),
        ("1.2", "Motivation", "4", False),
        ("1.3", "Objectives", "5", False),
        ("2", "DATASET", "7", True),
        ("2.1", "Dataset Source", "7", False),
        ("2.2", "Number of Records", "7", False),
        ("2.3", "Attributes", "8", False),
        ("2.4", "Target Variable", "9", False),
        ("3", "DATA PREPROCESSING", "10", True),
        ("3.1", "Missing Values", "10", False),
        ("3.2", "Cleaning", "10", False),
        ("3.3", "Transformation", "11", False),
        ("3.4", "Normalization & Partitioning", "12", False),
        ("4", "DATA MINING TECHNIQUES", "14", True),
        ("4.1", "Data Mining Process (KDD Framework)", "14", False),
        ("4.2", "Exploratory Data Mining & Feature Analysis", "15", False),
        ("4.3", "Supervised Classification & Detection Model", "16", False),
        ("4.4", "Pattern Discovery & Temporal Mining", "16", False),
        ("5", "IMPLEMENTATION", "18", True),
        ("5.1", "Python Development Environment & Tools", "18", False),
        ("5.2", "Implementation Pipeline & Architecture (KDD)", "18", False),
        ("5.3", "Parameters & Configurations", "20", False),
        ("6", "RESULTS", "21", True),
        ("6.1", "Tables", "21", False),
        ("6.2", "Graphs", "22", False),
        ("6.3", "Evaluation Metrics Analysis", "23", False),
        ("7", "COMPARISON AND DISCUSSION", "24", True),
        ("7.1", "Algorithm Comparison", "24", False),
        ("7.2", "Advantages", "25", False),
        ("7.3", "Limitations", "25", False),
        ("8", "CONCLUSION AND FUTURE SCOPE", "26", True),
        ("8.1", "Conclusion", "26", False),
        ("8.2", "Future Scope", "26", False),
        ("", "REFERENCES", "27", True),
    ]

    # Create table with 1 header row + len(toc_data) rows, 3 columns
    table = doc.add_table(rows=len(toc_data) + 1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Widths: Col 0: 1.4 in, Col 1: 3.8 in, Col 2: 0.8 in (Total = 6.0 in)
    col_widths = [Inches(1.4), Inches(3.8), Inches(0.8)]

    # Borderless styling
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)

    # Style Header Row
    hdr_cells = table.rows[0].cells
    headers = ["CHAPTER NO.", "TITLE", "PAGE NO."]
    hdr_alignments = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT]

    for i, h_text in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        p = cell.paragraphs[0]
        p.alignment = hdr_alignments[i]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(h_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    # Style Data Rows
    for r_idx, (chap_no, title, page_no, is_main) in enumerate(toc_data):
        row_cells = table.rows[r_idx + 1].cells

        # Col 0: CHAPTER NO.
        c0 = row_cells[0]
        c0.width = col_widths[0]
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(2.0 if is_main else 1.0)
        p0.paragraph_format.space_after = Pt(2.0 if is_main else 1.0)
        p0.paragraph_format.line_spacing = 1.15
        if chap_no:
            r0 = p0.add_run(chap_no)
            r0.font.name = "Times New Roman"
            r0.font.size = Pt(11)
            r0.bold = is_main
            r0.font.color.rgb = RGBColor(0, 0, 0)

        # Col 1: TITLE
        c1 = row_cells[1]
        c1.width = col_widths[1]
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(2.0 if is_main else 1.0)
        p1.paragraph_format.space_after = Pt(2.0 if is_main else 1.0)
        p1.paragraph_format.line_spacing = 1.15
        r1 = p1.add_run(title)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(11)
        r1.bold = is_main
        r1.font.color.rgb = RGBColor(0, 0, 0)

        # Col 2: PAGE NO.
        c2 = row_cells[2]
        c2.width = col_widths[2]
        p2 = c2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.paragraph_format.space_before = Pt(2.0 if is_main else 1.0)
        p2.paragraph_format.space_after = Pt(2.0 if is_main else 1.0)
        p2.paragraph_format.line_spacing = 1.15
        r2 = p2.add_run(page_no)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(11)
        r2.bold = is_main
        r2.font.color.rgb = RGBColor(0, 0, 0)

    # Set explicit column width on every cell
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    # Move table XML right after p_toc_heading
    p_toc_heading._p.addnext(table._tbl)

    # Find the paragraphs between p_toc_heading and ABSTRACT heading
    # Remove all the old broken paragraphs
    abstract_p = None
    for p in doc.paragraphs[toc_heading_idx + 1:]:
        if p.text.strip().upper() == 'ABSTRACT' and any(r.bold for r in p.runs):
            abstract_p = p
            break
    
    print('Found ABSTRACT heading paragraph:', abstract_p.text)

    # Remove all paragraphs between p_toc_heading and abstract_p
    current_p = p_toc_heading._p.getnext()
    removed_count = 0
    while current_p is not None and current_p != abstract_p._p:
        next_p = current_p.getnext()
        # Only remove if it's a paragraph (<w:p>), not our newly added table (<w:tbl>)
        if current_p.tag.endswith('p'):
            current_p.getparent().remove(current_p)
            removed_count += 1
        current_p = next_p

    print(f'Removed {removed_count} old broken TOC and spacing paragraphs.')

    # Ensure ABSTRACT starts cleanly on a fresh page
    abstract_p.paragraph_format.page_break_before = True

    # Save to fixed docx
    doc.save('doc/DataMining_Report_fixed.docx')
    print('Successfully saved to doc/DataMining_Report_fixed.docx!')

    # Try saving to DataMining_Report.docx as well
    try:
        doc.save('doc/DataMining_Report.docx')
        print('Successfully saved to doc/DataMining_Report.docx!')
    except Exception as e:
        print('doc/DataMining_Report.docx locked by Word:', e)

if __name__ == '__main__':
    build_perfect_report()
