import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def update_text_frame_content(tf, paragraphs_data, default_font_size=None, default_font_name="Calibri"):
    """
    Safely updates text_frame content while preserving structure.
    paragraphs_data can be:
      - list of strings
      - list of (bold_prefix, normal_body) tuples
    """
    # Clear text in first paragraph
    p0 = tf.paragraphs[0]
    p0.text = ""
    
    # Remove extra paragraphs from previous content
    p_elements = list(tf._txBody.p_lst)
    for p_elem in p_elements[1:]:
        tf._txBody.remove(p_elem)
        
    for idx, item in enumerate(paragraphs_data):
        if idx == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
            
        if isinstance(item, tuple):
            bold_prefix, normal_body = item
            r1 = p.add_run()
            r1.text = bold_prefix
            r1.font.bold = True
            if default_font_name:
                r1.font.name = default_font_name
            if default_font_size:
                r1.font.size = Pt(default_font_size)
                
            r2 = p.add_run()
            r2.text = normal_body
            r2.font.bold = False
            if default_font_name:
                r2.font.name = default_font_name
            if default_font_size:
                r2.font.size = Pt(default_font_size)
        else:
            r = p.add_run()
            r.text = str(item)
            if default_font_name:
                r.font.name = default_font_name
            if default_font_size:
                r.font.size = Pt(default_font_size)
                
        p.space_after = Pt(6)

print("Helper function defined.")
