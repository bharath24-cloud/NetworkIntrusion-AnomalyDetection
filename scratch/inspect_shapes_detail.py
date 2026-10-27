import pptx

prs = pptx.Presentation('doc/NEXUS_Project_PPT.pptx')
print(f"Slide dimensions: {prs.slide_width / 914400} x {prs.slide_height / 914400} inches")

for i, slide in enumerate(prs.slides):
    print(f"Slide {i+1}: layout={slide.slide_layout.name}, shapes={len(slide.shapes)}")
    for j, shape in enumerate(slide.shapes):
        t = ""
        if shape.has_text_frame:
            t = shape.text_frame.text.replace("\n", " ")[:60]
        elif shape.has_table:
            t = f"Table {len(shape.table.rows)}x{len(shape.table.columns)}"
        elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
            t = "Picture"
        print(f"  Shape {j}: name='{shape.name}', type={shape.shape_type}, text='{t}'")
