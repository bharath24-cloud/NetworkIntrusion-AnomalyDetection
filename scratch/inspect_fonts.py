import pptx

prs = pptx.Presentation('doc/NEXUS_Project_PPT.pptx')
slide = prs.slides[2] # Slide 3 (Introduction)
for s in slide.shapes:
    if s.has_text_frame:
        print(f"Shape: {s.name}")
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                print(f"  Run text: {repr(r.text[:30])}, font={r.font.name}, size={r.font.size}, bold={r.font.bold}, color={r.font.color.type if r.font.color else None}")
